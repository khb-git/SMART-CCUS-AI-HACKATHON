import json


def test_xlsx_ingestion_extracts_sheet_as_table_chunk(tmp_path):
    from openpyxl import Workbook
    import ingestion.main as ingestion_main

    raw_dir = tmp_path / "raw_docs"
    raw_dir.mkdir()

    xlsx_path = raw_dir / "completeness_tool.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Checklist"
    sheet.append(["Requirement", "CFR", "Status"])
    sheet.append(["Site map", "146.82(a)(2)", "Required"])
    sheet.append(["AoR model", "146.84", "Required"])
    workbook.save(xlsx_path)

    output_dir = tmp_path / "chunked"

    count = ingestion_main.process_xlsx(
        file_path=xlsx_path,
        output_root=output_dir,
        metadata={
            "url": "https://example.com/completeness_tool.xlsx",
            "source_page": "https://example.com/docket",
            "summary": "Class VI completeness tool",
        },
    )

    assert count == 1

    document_dir = output_dir / "completeness_tool"
    assert document_dir.exists()

    document_attr = json.loads(
        (document_dir / "attribute.json").read_text(encoding="utf-8")
    )

    assert document_attr["datasource_name"] == "completeness_tool.xlsx"
    assert document_attr["file_type"] == "XLSX"
    assert document_attr["online_link"] == "https://example.com/completeness_tool.xlsx"

    chunk_text = (document_dir / "0" / "content.txt").read_text(encoding="utf-8")
    chunk_attr = json.loads(
        (document_dir / "0" / "attribute.json").read_text(encoding="utf-8")
    )

    assert "| Requirement | CFR | Status |" in chunk_text
    assert "| Site map | 146.82(a)(2) | Required |" in chunk_text
    assert chunk_attr["content_type"] == "table"
    assert chunk_attr["sheet_name"] == "Checklist"
    assert chunk_attr["row_start"] == 1
    assert chunk_attr["row_end"] == 3


def test_xlsx_ingestion_splits_large_sheet_by_row_windows(tmp_path):
    from openpyxl import Workbook
    import ingestion.main as ingestion_main

    xlsx_path = tmp_path / "large_sheet.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Large"
    sheet.append(["Col A", "Col B"])

    for i in range(1, 6):
        sheet.append([f"A{i}", f"B{i}"])

    workbook.save(xlsx_path)

    documents = ingestion_main.extract_xlsx_documents(
        file_path=xlsx_path,
        max_rows_per_chunk=3,
    )

    assert len(documents) == 2

    assert documents[0].metadata["sheet_name"] == "Large"
    assert documents[0].metadata["row_start"] == 1
    assert documents[0].metadata["row_end"] == 3

    assert documents[1].metadata["row_start"] == 4
    assert documents[1].metadata["row_end"] == 6


def test_chunk_from_manifest_processes_xlsx(tmp_path):
    from openpyxl import Workbook
    import ingestion.main as ingestion_main

    raw_dir = tmp_path / "raw_docs"
    raw_dir.mkdir()

    xlsx_path = raw_dir / "class_vi_tool.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tool"
    sheet.append(["Field", "Description"])
    sheet.append(["Injection Zone", "Named formation"])
    workbook.save(xlsx_path)

    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(
        json.dumps([
            {
                "summary": "Completeness tool",
                "url": "https://example.com/class_vi_tool.xlsx",
                "source_page": "https://example.com/docket",
                "local_path": str(xlsx_path),
            }
        ]),
        encoding="utf-8",
    )

    output_dir = tmp_path / "chunked"

    stats = ingestion_main.chunk_from_manifest(
        manifest_path=manifest_path,
        output_root=output_dir,
    )

    assert stats["processed"] == 1
    assert stats["skipped_unsupported_type"] == 0

    document_dir = output_dir / "class_vi_tool"
    assert document_dir.exists()

    chunk_attr = json.loads(
        (document_dir / "0" / "attribute.json").read_text(encoding="utf-8")
    )

    assert chunk_attr["content_type"] == "table"
    assert chunk_attr["sheet_name"] == "Tool"