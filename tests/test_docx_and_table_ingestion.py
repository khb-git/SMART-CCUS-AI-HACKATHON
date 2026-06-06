import json


def test_docx_ingestion_extracts_paragraphs_and_tables(tmp_path):
    from docx import Document
    import ingestion.main as ingestion_main

    raw_dir = tmp_path / "raw_docs"
    raw_dir.mkdir()

    docx_path = raw_dir / "template.docx"

    doc = Document()
    doc.add_paragraph("This is the application narrative template.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Description"
    table.cell(1, 0).text = "Injection Zone"
    table.cell(1, 1).text = "Named geologic formation"

    doc.save(docx_path)

    output_dir = tmp_path / "chunked"

    count = ingestion_main.process_docx(
        file_path=docx_path,
        output_root=output_dir,
        metadata={
            "url": "https://example.com/template.docx",
            "source_page": "https://example.com/docket",
            "summary": "Application narrative template",
        },
    )

    assert count >= 2

    document_dir = output_dir / "template"
    assert document_dir.exists()

    document_attr = json.loads(
        (document_dir / "attribute.json").read_text(encoding="utf-8")
    )

    assert document_attr["datasource_name"] == "template.docx"
    assert document_attr["file_type"] == "DOCX"
    assert document_attr["online_link"] == "https://example.com/template.docx"

    chunk_types = []

    for child in document_dir.iterdir():
        if not child.is_dir():
            continue
        attr_path = child / "attribute.json"
        if attr_path.exists():
            chunk_attr = json.loads(attr_path.read_text(encoding="utf-8"))
            chunk_types.append(chunk_attr["content_type"])

    assert "text" in chunk_types
    assert "table" in chunk_types


def test_pdf_table_documents_are_written_as_table_chunks(tmp_path, monkeypatch):
    import ingestion.main as ingestion_main

    pdf_path = tmp_path / "permit.pdf"
    pdf_path.write_bytes(b"%PDF fake content")

    output_dir = tmp_path / "chunked"

    class FakeLoader:
        def __init__(self, path):
            self.path = path

        def load(self):
            return [
                ingestion_main.IngestionDocument(
                    page_content="Normal PDF text.",
                    metadata={"page": 0},
                )
            ]

    class FakeSplitter:
        def __init__(self, chunk_size=1000, chunk_overlap=100, add_start_index=True):
            pass

        def split_documents(self, pages):
            return pages

    monkeypatch.setattr(ingestion_main, "PyPDFLoader", FakeLoader)
    monkeypatch.setattr(ingestion_main, "RecursiveCharacterTextSplitter", FakeSplitter)
    monkeypatch.setattr(
        ingestion_main,
        "extract_pdf_table_documents",
        lambda file_path: [
            ingestion_main.IngestionDocument(
                page_content="| Parameter | Value |\n| --- | --- |\n| Porosity | 0.20 |",
                metadata={
                    "page": 0,
                    "content_type": "table",
                    "table_index": 0,
                },
            )
        ],
    )

    count = ingestion_main.process_pdf(
        file_path=pdf_path,
        output_root=output_dir,
        metadata={
            "url": "https://example.com/permit.pdf",
            "source_page": "https://example.com/docket",
            "summary": "Permit PDF",
        },
    )

    assert count == 2

    document_dir = output_dir / "permit"

    chunk_types = []
    table_found = False

    for child in document_dir.iterdir():
        if not child.is_dir():
            continue

        attr = json.loads((child / "attribute.json").read_text(encoding="utf-8"))
        text = (child / "content.txt").read_text(encoding="utf-8")

        chunk_types.append(attr["content_type"])

        if attr["content_type"] == "table":
            table_found = True
            assert attr["table_index"] == 0
            assert "| Porosity | 0.20 |" in text

    assert "text" in chunk_types
    assert table_found


def test_chunk_from_manifest_now_processes_docx(tmp_path):
    from docx import Document
    import ingestion.main as ingestion_main

    raw_dir = tmp_path / "raw_docs"
    raw_dir.mkdir()

    docx_path = raw_dir / "application_template.docx"

    doc = Document()
    doc.add_paragraph("Template paragraph for Class VI application.")
    doc.save(docx_path)

    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(
        json.dumps([
            {
                "summary": "Application template",
                "url": "https://example.com/application_template.docx",
                "source_page": "https://example.com/docket",
                "local_path": str(docx_path),
            }
        ]),
        encoding="utf-8",
    )

    output_dir = tmp_path / "chunked"

    stats = ingestion_main.chunk_from_manifest(
        manifest_path=manifest_path,
        output_root=output_dir,
    )

    assert stats["processed"] == 1
    assert stats["skipped_unsupported_type"] == 0
    assert (output_dir / "application_template").exists()