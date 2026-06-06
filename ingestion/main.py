import argparse
import json
import os
import uuid
from pathlib import Path

from dotenv import load_dotenv
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

from rag.manifest import load_manifest

from dataclasses import dataclass

import re

load_dotenv()

# Environment variables
source_dir = os.getenv("PATH_TO_RAW_DATA_DIR")
ingestion_folder = os.getenv("PATH_TO_CHUNKED_DATA_DIR")
manifest_path = os.getenv("PATH_TO_MANIFEST")

# Constant for EPA source
EPA_LINK = "https://www.epa.gov/uic/final-class-vi-guidance-documents"

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}

REVIEW_SECTION_METADATA = {
    "project_narrative": {
        "schema_section_id": "1",
        "schema_section_title": "Project Narrative",
    },
    "site_geologic_characterization": {
        "schema_section_id": "2",
        "schema_section_title": "Site Geologic Characterization",
    },
    "aor_corrective_action": {
        "schema_section_id": "3",
        "schema_section_title": "AoR and Corrective Action Plan",
    },
    "financial_responsibility": {
        "schema_section_id": "4",
        "schema_section_title": "Financial Responsibility",
    },
    "well_construction": {
        "schema_section_id": "5",
        "schema_section_title": "Well Construction Plan",
    },
    "pre_operational_testing": {
        "schema_section_id": "6",
        "schema_section_title": "Pre-Operational Testing Plan",
    },
    "site_operating": {
        "schema_section_id": "7",
        "schema_section_title": "Site Operating Plan",
    },
    "testing_monitoring": {
        "schema_section_id": "8",
        "schema_section_title": "Testing and Monitoring Plan",
    },
    "injection_well_plugging": {
        "schema_section_id": "9",
        "schema_section_title": "Injection Well Plugging Plan",
    },
    "pisc_site_closure": {
        "schema_section_id": "10",
        "schema_section_title": "PISC and Site Closure Plan",
    },
    "emergency_remedial_response": {
        "schema_section_id": "11",
        "schema_section_title": "Emergency and Remedial Response Plan",
    },
    "unknown": {
        "schema_section_id": "",
        "schema_section_title": "",
    },
}


def normalize_for_matching(value):
    """Normalize text for lightweight plan-type matching."""
    return (
        str(value or "")
        .lower()
        .replace("-", " ")
        .replace("_", " ")
        .replace("+", " ")
    )


def infer_plan_type(file_path, metadata=None):
    """Infer the Class VI plan/review type from filename and manifest metadata.

    This is intentionally lightweight and rules-based for now. It gives the
    future vector store useful metadata before we implement deeper schema-aware
    section detection.
    """
    metadata = metadata or {}
    file_path = Path(file_path)

    combined = " ".join(
        [
            normalize_for_matching(file_path.name),
            normalize_for_matching(metadata.get("summary", "")),
            normalize_for_matching(metadata.get("url", "")),
            normalize_for_matching(metadata.get("source_page", "")),
        ]
    )

    # Order matters: more specific plans before broader narrative matches.
    if any(term in combined for term in ["testing and monitoring", "tm plan", "t m plan"]):
        return "testing_monitoring"

    if any(term in combined for term in ["pre operational", "pre operation", "preoperational"]):
        return "pre_operational_testing"

    if any(term in combined for term in ["well construction", "construction details"]):
        return "well_construction"

    if any(term in combined for term in ["plugging plan", "injection well plugging"]):
        return "injection_well_plugging"

    if any(term in combined for term in ["emergency and remedial", "errp", "err plan"]):
        return "emergency_remedial_response"

    if any(term in combined for term in ["pisc", "post injection", "site closure"]):
        return "pisc_site_closure"

    if any(term in combined for term in ["financial responsibility", "cost estimates", "fr demonstration"]):
        return "financial_responsibility"

    if any(term in combined for term in ["aor", "area of review", "corrective action"]):
        return "aor_corrective_action"

    if any(term in combined for term in ["site operating", "operating plan", "operations plan"]):
        return "site_operating"

    if any(term in combined for term in ["site geologic", "geologic characterization", "site characterization"]):
        return "site_geologic_characterization"

    if any(term in combined for term in ["project narrative", "application narrative", "narrative"]):
        return "project_narrative"

    return "unknown"


def build_review_metadata(file_path, metadata=None):
    """Build review-schema metadata for a document/chunk."""
    plan_type = infer_plan_type(file_path, metadata)
    section_metadata = REVIEW_SECTION_METADATA.get(
        plan_type,
        REVIEW_SECTION_METADATA["unknown"],
    )

    return {
        "plan_type": plan_type,
        "schema_section_id": section_metadata["schema_section_id"],
        "schema_section_title": section_metadata["schema_section_title"],
    }

def build_splitter(chunk_size=1000, chunk_overlap=100):
    """Create the LangChain text splitter used by ingestion."""
    return RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        add_start_index=True,
    )

SECTION_HEADING_PATTERNS = [
    r"^\s*(\d+(?:\.\d+){0,4})\s+(.{3,120})$",
    r"^\s*([A-Z]\.\d+(?:\.\d+)*)\s+(.{3,120})$",
    r"^\s*(Appendix\s+[A-Z0-9]+[:.\-\s]+.{3,120})$",
    r"^\s*(Section\s+\d+(?:\.\d+){0,4}[:.\-\s]+.{3,120})$",
]


def normalize_heading_text(value):
    """Normalize heading text for metadata and chunk context."""
    return " ".join(str(value or "").split()).strip()


def looks_like_section_heading(line):
    """Return True if a line looks like a section/subsection heading."""
    line = normalize_heading_text(line)

    if not line:
        return False

    if len(line) > 140:
        return False

    for pattern in SECTION_HEADING_PATTERNS:
        if re.match(pattern, line, flags=re.IGNORECASE):
            return True

    # Common technical section headings without numbering.
    lowered = line.lower()
    heading_keywords = [
        "injection rate and pressure monitoring",
        "continuous recording",
        "testing and monitoring",
        "groundwater monitoring",
        "plume and pressure front tracking",
        "mechanical integrity",
        "well construction",
        "area of review",
        "corrective action",
        "site closure",
        "emergency and remedial response",
    ]

    return any(keyword in lowered for keyword in heading_keywords)


def extract_section_heading(line):
    """Extract a stable heading string from a possible heading line."""
    line = normalize_heading_text(line)

    if not line:
        return ""

    for pattern in SECTION_HEADING_PATTERNS:
        match = re.match(pattern, line, flags=re.IGNORECASE)
        if match:
            return normalize_heading_text(line)

    if looks_like_section_heading(line):
        return line

    return ""


def annotate_documents_with_section_context(documents):
    """Attach nearest detected heading to each document before splitting.

    This is a lightweight section-aware layer. It does not fully parse a PDF
    outline, but it preserves useful local heading context for retrieval.
    """
    current_heading = ""

    for doc in documents:
        text = getattr(doc, "page_content", "") or ""
        first_heading = ""

        for line in text.splitlines():
            heading = extract_section_heading(line)
            if heading:
                first_heading = heading
                current_heading = heading
                break

        if current_heading:
            doc.metadata["section_heading"] = current_heading
            doc.metadata["local_section_title"] = current_heading

        if first_heading:
            doc.metadata["detected_heading_on_page"] = first_heading

    return documents


def add_section_context_to_chunks(chunks):
    """Add section metadata and prepend heading context to text chunks."""
    for chunk in chunks:
        heading = normalize_heading_text(
            chunk.metadata.get("section_heading")
            or chunk.metadata.get("local_section_title")
            or ""
        )

        if heading:
            chunk.metadata["section_heading"] = heading
            chunk.metadata["local_section_title"] = heading

            prefix = f"Section context: {heading}\n\n"
            if not chunk.page_content.startswith(prefix):
                chunk.page_content = prefix + chunk.page_content

    return chunks

def write_chunks_for_document(file_path, chunks, output_root, metadata=None):
    """Write one document's chunks into the existing folder structure.

    Output shape:
        output_root/
            document_stem/
                attribute.json
                0/content.txt
                0/attribute.json
                1/content.txt
                1/attribute.json
    """
    metadata = metadata or {}
    review_metadata = build_review_metadata(file_path, metadata)

    file_path = Path(file_path)
    output_root = Path(output_root)
    output_root.mkdir(parents=True, exist_ok=True)

    file_folder_name = file_path.stem
    file_output_dir = output_root / file_folder_name
    file_output_dir.mkdir(parents=True, exist_ok=True)

    doc_uuid = str(uuid.uuid4())

    datasource_attr = {
        "document_id": doc_uuid,
        "datasource_name": file_path.name,
        "file_type": file_path.suffix.replace(".", "").upper(),
        "local_path": str(file_path),
        "online_link": metadata.get("url", EPA_LINK),
        "source_page": metadata.get("source_page", ""),
        "summary": metadata.get("summary", ""),
        "author_name": metadata.get("author_name", "Environmental Protection Agency"),
        "plan_type": review_metadata["plan_type"],
        "schema_section_id": review_metadata["schema_section_id"],
        "schema_section_title": review_metadata["schema_section_title"],
    }

    with (file_output_dir / "attribute.json").open("w", encoding="utf-8") as f:
        json.dump(datasource_attr, f, indent=4)

    for i, chunk in enumerate(chunks):
        chunk_folder = file_output_dir / str(i)
        chunk_folder.mkdir(parents=True, exist_ok=True)

        chunk_uuid = str(uuid.uuid4())

        with (chunk_folder / "content.txt").open("w", encoding="utf-8") as f:
            f.write(chunk.page_content)

        content_type = chunk.metadata.get("content_type", "text")

        chunk_attr = {
            "chunk_id": chunk_uuid,
            "parent_document_id": doc_uuid,
            "page": chunk.metadata.get("page", 0) + 1,
            "chunk_index": i,
            "content_type": content_type,
            "table_index": chunk.metadata.get("table_index"),
            "sheet_name": chunk.metadata.get("sheet_name"),
            "row_start": chunk.metadata.get("row_start"),
            "row_end": chunk.metadata.get("row_end"),
            "datasource_name": file_path.name,
            "local_path": str(file_path),
            "online_link": metadata.get("url", EPA_LINK),
            "source_page": metadata.get("source_page", ""),
            "summary": metadata.get("summary", ""),
            "plan_type": review_metadata["plan_type"],
            "schema_section_id": review_metadata["schema_section_id"],
            "schema_section_title": review_metadata["schema_section_title"],
            "section_heading": chunk.metadata.get("section_heading", ""),
            "local_section_title": chunk.metadata.get("local_section_title", ""),
            "detected_heading_on_page": chunk.metadata.get("detected_heading_on_page", ""),
        }

        with (chunk_folder / "attribute.json").open("w", encoding="utf-8") as f:
            json.dump(chunk_attr, f, indent=4)

    return len(chunks)

@dataclass
class IngestionDocument:
    """Small document object compatible with the existing chunk writer."""
    page_content: str
    metadata: dict


def table_to_markdown(rows):
    """Convert a list of table rows into a simple Markdown table."""
    cleaned_rows = [
        [str(cell or "").replace("\n", " ").strip() for cell in row]
        for row in rows
        if any(str(cell or "").strip() for cell in row)
    ]

    if not cleaned_rows:
        return ""

    max_cols = max(len(row) for row in cleaned_rows)
    normalized = [
        row + [""] * (max_cols - len(row))
        for row in cleaned_rows
    ]

    header = normalized[0]
    separator = ["---"] * max_cols
    body = normalized[1:]

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]

    for row in body:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def extract_pdf_table_documents(file_path):
    """Extract PDF tables as table chunks using pdfplumber.

    These are separate from normal PDF text chunks so technical values in
    tables remain searchable downstream.
    """
    import pdfplumber

    table_docs = []

    with pdfplumber.open(str(file_path)) as pdf:
        for page_index, page in enumerate(pdf.pages):
            tables = page.extract_tables() or []

            for table_index, table in enumerate(tables):
                markdown = table_to_markdown(table)
                if not markdown:
                    continue

                table_docs.append(
                    IngestionDocument(
                        page_content=markdown,
                        metadata={
                            "page": page_index,
                            "content_type": "table",
                            "table_index": table_index,
                        },
                    )
                )

    return table_docs


def extract_docx_documents(file_path):
    """Extract DOCX paragraphs and tables as ingestion documents."""
    from docx import Document

    doc = Document(str(file_path))
    documents = []

    paragraph_text = "\n".join(
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph.text.strip()
    )

    if paragraph_text:
        documents.append(
            IngestionDocument(
                page_content=paragraph_text,
                metadata={
                    "page": 0,
                    "content_type": "text",
                },
            )
        )

    for table_index, table in enumerate(doc.tables):
        rows = [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ]

        markdown = table_to_markdown(rows)
        if not markdown:
            continue

        documents.append(
            IngestionDocument(
                page_content=markdown,
                metadata={
                    "page": 0,
                    "content_type": "table",
                    "table_index": table_index,
                },
            )
        )

    return documents


def extract_xlsx_documents(file_path, max_rows_per_chunk=50):
    """Extract XLSX sheets as Markdown table chunks.

    Each sheet is chunked by row windows so large spreadsheets do not become
    one huge chunk. XLSX content is treated as table data only.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(file_path), data_only=True)
    documents = []

    for sheet in workbook.worksheets:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            cleaned_row = [
                "" if cell is None else str(cell).replace("\n", " ").strip()
                for cell in row
            ]

            if any(cell for cell in cleaned_row):
                rows.append(cleaned_row)

        if not rows:
            continue

        for start in range(0, len(rows), max_rows_per_chunk):
            window = rows[start:start + max_rows_per_chunk]
            markdown = table_to_markdown(window)

            if not markdown:
                continue

            documents.append(
                IngestionDocument(
                    page_content=markdown,
                    metadata={
                        "page": 0,
                        "content_type": "table",
                        "sheet_name": sheet.title,
                        "row_start": start + 1,
                        "row_end": start + len(window),
                    },
                )
            )

    return documents

def process_pdf(file_path, output_root, metadata=None, chunk_size=1000, chunk_overlap=100):
    """Load one PDF, split text, extract tables, and write chunk files."""
    file_path = Path(file_path)

    loader = PyPDFLoader(str(file_path))
    pages = loader.load()
    pages = annotate_documents_with_section_context(pages)

    splitter = build_splitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    text_chunks = splitter.split_documents(pages)
    text_chunks = add_section_context_to_chunks(text_chunks)

    for chunk in text_chunks:
        chunk.metadata["content_type"] = "text"

    try:
        table_chunks = extract_pdf_table_documents(file_path)
    except Exception as exc:
        print(f"Warning: failed to extract PDF tables from {file_path.name}: {exc}")
        table_chunks = []

    all_chunks = text_chunks + table_chunks

    return write_chunks_for_document(
        file_path=file_path,
        chunks=all_chunks,
        output_root=output_root,
        metadata=metadata,
    )


def process_docx(file_path, output_root, metadata=None, chunk_size=1000, chunk_overlap=100):
    """Load one DOCX, split paragraph text, extract tables, and write chunk files."""
    file_path = Path(file_path)

    documents = extract_docx_documents(file_path)

    splitter = build_splitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )

    final_chunks = []

    for doc in documents:
        if doc.metadata.get("content_type") == "table":
            final_chunks.append(doc)
        else:
            annotated_docs = annotate_documents_with_section_context([doc])
            split_docs = splitter.split_documents(annotated_docs)
            split_docs = add_section_context_to_chunks(split_docs)

            for chunk in split_docs:
                chunk.metadata["content_type"] = "text"

            final_chunks.extend(split_docs)

    return write_chunks_for_document(
        file_path=file_path,
        chunks=final_chunks,
        output_root=output_root,
        metadata=metadata,
    )

def process_xlsx(file_path, output_root, metadata=None, chunk_size=1000, chunk_overlap=100):
    """Load one XLSX workbook and write sheet/table chunks."""
    file_path = Path(file_path)

    table_chunks = extract_xlsx_documents(file_path)

    return write_chunks_for_document(
        file_path=file_path,
        chunks=table_chunks,
        output_root=output_root,
        metadata=metadata,
    )

def process_file(file_path, output_root, metadata=None, chunk_size=1000, chunk_overlap=100):
    """Dispatch supported files to the correct LangChain/table-aware processor."""
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return process_pdf(
            file_path=file_path,
            output_root=output_root,
            metadata=metadata,
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )

    if suffix == ".docx":
        return process_docx(
            file_path=file_path,
            output_root=output_root,
            metadata=metadata,
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )
    if suffix == ".xlsx":
        return process_xlsx(
            file_path=file_path,
            output_root=output_root,
            metadata=metadata,
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )

    raise ValueError(f"Unsupported file type: {suffix}")

def chunk_with_langchain(source_dir, output_root, chunk_size=1000, chunk_overlap=100):
    """Process all supported files in a raw data directory."""
    source_dir = Path(source_dir)
    output_root = Path(output_root)

    if not source_dir.exists():
        print(f"Source directory not found: {source_dir}")
        return {
            "processed": 0,
            "skipped": 0,
            "failed": 0,
        }

    stats = {
        "processed": 0,
        "skipped": 0,
        "failed": 0,
    }

    for file_path in source_dir.iterdir():
        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            stats["skipped"] += 1
            continue

        try:
            count = process_file(
                file_path=file_path,
                output_root=output_root,
                metadata={
                    "url": EPA_LINK,
                    "source_page": "",
                    "summary": "",
                },
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
            )
            stats["processed"] += 1
            print(f"Successfully processed {count} chunks for: {file_path.name}")

        except Exception as e:
            stats["failed"] += 1
            print(f"Failed to process {file_path.name}: {str(e)}")

    return stats


def chunk_from_manifest(manifest_path, output_root, chunk_size=1000, chunk_overlap=100):
    """Process downloaded files listed in a scraper manifest.

    Uses each manifest entry's local_path field, which is written by
    rag.manifest.download_manifest().
    """
    entries = load_manifest(manifest_path)

    stats = {
        "processed": 0,
        "skipped_missing_local_path": 0,
        "skipped_missing_file": 0,
        "skipped_unsupported_type": 0,
        "failed": 0,
    }

    for entry in entries:
        if not entry.local_path:
            stats["skipped_missing_local_path"] += 1
            print(f"Skipping manifest entry with no local_path: {entry.url}")
            continue

        file_path = Path(entry.local_path)

        if not file_path.exists():
            stats["skipped_missing_file"] += 1
            print(f"Skipping missing local file: {file_path}")
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            stats["skipped_unsupported_type"] += 1
            print(f"Skipping unsupported file type: {file_path}")
            continue

        try:
            count = process_file(
                file_path=file_path,
                output_root=output_root,
                metadata={
                    "url": entry.url,
                    "source_page": entry.source_page,
                    "summary": entry.summary,
                },
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
            )
            stats["processed"] += 1
            print(f"Successfully processed {count} chunks for: {file_path.name}")

        except Exception as e:
            stats["failed"] += 1
            print(f"Failed to process {file_path.name}: {str(e)}")

    return stats


def parse_args():
    parser = argparse.ArgumentParser(
        description="LangChain ingestion for downloaded Class VI permit documents.",
    )
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=None,
        help="Directory of raw downloaded files. Used when --manifest is not provided.",
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        default=None,
        help="Manifest JSON with local_path values from the downloader.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Directory where chunked ingestion output should be written.",
    )
    parser.add_argument(
        "--chunk-size",
        type=int,
        default=1000,
        help="LangChain splitter chunk size. Default: 1000.",
    )
    parser.add_argument(
        "--chunk-overlap",
        type=int,
        default=100,
        help="LangChain splitter chunk overlap. Default: 100.",
    )
    return parser.parse_args()


def main():
    args = parse_args()

    selected_manifest = args.manifest or manifest_path
    selected_source_dir = args.source_dir or source_dir
    selected_output = args.output or ingestion_folder

    if not selected_output:
        raise ValueError(
            "No output directory provided. Use --output or PATH_TO_CHUNKED_DATA_DIR."
        )

    if selected_manifest:
        stats = chunk_from_manifest(
            manifest_path=selected_manifest,
            output_root=selected_output,
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
        )
    elif selected_source_dir:
        stats = chunk_with_langchain(
            source_dir=selected_source_dir,
            output_root=selected_output,
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
        )
    else:
        raise ValueError(
            "No input provided. Use --manifest, --source-dir, PATH_TO_MANIFEST, "
            "or PATH_TO_RAW_DATA_DIR."
        )

    print("Ingestion complete.")
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    main()